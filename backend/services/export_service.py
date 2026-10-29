import io
import csv
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.models.tender import Tender

class ExportService:
    @staticmethod
    def generate_excel(tenders: List[Tender]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "BitVivid Tender Pipeline"
        ws.sheet_view.showGridLines = True

        # Styles
        navy = PatternFill("solid", fgColor="17365D")
        blue = PatternFill("solid", fgColor="1F4E78")
        light_blue = PatternFill("solid", fgColor="D9EAF7")
        zebra = PatternFill("solid", fgColor="F7FAFD")
        white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0")
        )

        # Title Banner
        ws.merge_cells("A1:N1")
        ws["A1"] = "BITVIVID TENDER TRACKER & OPPORTUNITY PIPELINE"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = navy
        ws["A1"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        headers = [
            "Sr No", "Status", "Priority", "Tender ID", "Portal", "Tender Full Title",
            "Organisation", "Closing Date", "EMD", "Turnover",
            "Eligibility Fit", "Fit Score", "Direct Tender Link", "Internal Notes"
        ]

        # Header Row
        ws.row_dimensions[3].height = 28
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = white_font
            cell.fill = blue
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Data Rows
        for r_idx, t in enumerate(tenders, start=4):
            sr = r_idx - 3
            ws.row_dimensions[r_idx].height = 40
            row_data = [
                sr,
                t.status or "New",
                t.priority or "Normal",
                t.tender_id,
                t.portal,
                t.title,
                t.organisation or "",
                t.closing_date or "",
                t.emd_amount or "",
                t.turnover_required or "",
                t.eligibility_assessment or "",
                f"{t.fit_score}%",
                t.url or "",
                t.user_notes or ""
            ]
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

                if sr % 2 == 0:
                    cell.fill = zebra

                # Status coloring
                if c_idx == 2:
                    if val == "Bid Ready":
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")
                        cell.font = Font(color="006100", bold=True)
                    elif val == "Reviewing":
                        cell.fill = PatternFill("solid", fgColor="FFEB9C")
                        cell.font = Font(color="9C6500", bold=True)

                # Link styling
                if c_idx == 13 and val:
                    cell.hyperlink = val
                    cell.font = Font(color="0563C1", underline="single")

        # Column widths
        widths = {1: 8, 2: 14, 3: 12, 4: 22, 5: 18, 6: 45, 7: 30, 8: 20, 9: 18, 10: 18, 11: 30, 12: 12, 13: 35, 14: 30}
        for col_idx, width in widths.items():
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_csv(tenders: List[Tender]) -> io.StringIO:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Sr No", "Tender ID", "Status", "Priority", "Portal", "Title",
            "Organisation", "Closing Date", "EMD", "Turnover",
            "Eligibility Fit", "Fit Score", "Link", "User Notes"
        ])
        for idx, t in enumerate(tenders, 1):
            writer.writerow([
                idx, t.tender_id, t.status, t.priority, t.portal, t.title,
                t.organisation, t.closing_date, t.emd_amount, t.turnover_required,
                t.eligibility_assessment, f"{t.fit_score}%", t.url, t.user_notes
            ])
        output.seek(0)
        return output

